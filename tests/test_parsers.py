from __future__ import annotations

import csv
import io

from docx import Document
from openpyxl import Workbook

from evidenceops.parsers import extract_questions, parse_document


def _minimal_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(pdf)


def test_parse_txt_and_extract_questions() -> None:
    chunks = parse_document("questionnaire.md", b"# Security\nDo you encrypt data at rest?\nDescribe key rotation.\n")
    questions = extract_questions(chunks)
    assert [item[0] for item in questions] == ["Do you encrypt data at rest?", "Describe key rotation."]
    assert chunks[0].page_or_sheet.startswith("Lines")


def test_parse_xlsx_with_sheet_and_row_locator() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Controls"
    worksheet.append(["Control", "Evidence"])
    worksheet.append(["Encryption", "AES-256 is used for stored backups."])
    output = io.BytesIO()
    workbook.save(output)

    chunks = parse_document("evidence.xlsx", output.getvalue())

    assert chunks[0].page_or_sheet == "Sheet Controls, row 2"
    assert "Evidence: AES-256" in chunks[0].text


def test_parse_docx_and_csv() -> None:
    document = Document()
    document.add_paragraph("Backups are retained for 30 days.")
    output = io.BytesIO()
    document.save(output)
    assert "30 days" in parse_document("policy.docx", output.getvalue())[0].text

    csv_output = io.StringIO()
    writer = csv.writer(csv_output)
    writer.writerow(["Topic", "Statement"])
    writer.writerow(["MFA", "MFA is enabled for administrators."])
    chunks = parse_document("controls.csv", csv_output.getvalue().encode())
    assert chunks[0].page_or_sheet == "CSV row 2"


def test_parse_pdf_preserves_page_locator() -> None:
    chunks = parse_document("security.pdf", _minimal_pdf("Encryption at rest is enabled."))

    assert chunks[0].page_or_sheet == "Page 1"
    assert "Encryption at rest is enabled." in chunks[0].text
